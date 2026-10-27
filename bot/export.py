import io

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .i18n import district_name, t

HEAD = {
    "ru": ["Дата", "Сделка", "Тип", "Заголовок", "Цена", "Цена, $", "Комнаты", "Площадь", "Этаж", "Район", "Кто", "Ссылка"],
    "uz": ["Sana", "Bitim", "Turi", "Sarlavha", "Narx", "Narx, $", "Xonalar", "Maydon", "Qavat", "Tuman", "Kim", "Havola"],
}


def build_excel(items: list[dict], lang: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "OLX"
    ws.append(HEAD.get(lang, HEAD["ru"]))
    for i in items:
        ws.append([
            (i.get("created_at") or i["first_seen"])[:16].replace("T", " "),
            t(f"deal_{i['deal']}", lang),
            t(f"ptype_{i['ptype']}", lang),
            i["title"],
            i.get("price_text"),
            round(i["price_usd"]) if i.get("price_usd") else None,
            i.get("rooms"),
            i.get("area"),
            i.get("floor"),
            district_name(i["district"], lang) if i.get("district") else i.get("district_raw"),
            t("owner", lang) if i.get("is_owner") else t("agent", lang),
            i["url"],
        ])
    for col, width in enumerate([16, 10, 11, 45, 16, 9, 8, 9, 7, 18, 22, 50], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
